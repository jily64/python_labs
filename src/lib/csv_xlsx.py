import csv
import openpyxl
from src.lib.tables import check_rows

def csv_to_xlsx(csv_path: str, xlsx_path: str) -> None:
    """
    Функция конвертирует .csv в .xlsx
    Используется дополнительная библиотека openpyxl.
    В проверках сделано так, что программа требует наличие хеддера и как минмум 1 ряда с данными.
    
    Открываем csv, читаем семпл, полуаем тип разделителей (диалект), читаем уже полный файл, записываем в xlsx через библиотеку openpyxl.
    И меняем размеры колонок (берем максимальную длину значения в колонке + 3)
    """
    
    try:
        # xlsx converter setup
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"  # type: ignore

        # Reader
        with open(csv_path, "r", encoding="utf-8") as f:
            sample = f.read(2048)
            dialect = csv.Sniffer().sniff(sample)
            f.seek(0)
            
            reader = csv.reader(f, dialect=dialect)
            rows = list(reader)

        #Checkers
        if len(rows) <= 1:
            raise ValueError("Not enough values in .csv. File must have at least 2 rows (Header + 1 data row)")
        check_rows(rows)  # type: ignore

        # Writer
        for row in rows:
            ws.append(row)  # type: ignore
            
        for col in ws.columns:  # type: ignore
          max_len = max(len(str(cell.value or "")) for cell in col)
          col_letter = openpyxl.utils.get_column_letter(col[0].column)  # type: ignore
          ws.column_dimensions[col_letter].width = max(max_len + 3, 8)  # type: ignore
        
        wb.save(xlsx_path)

    # Errors proceeder
    except FileNotFoundError as e:
        raise FileNotFoundError("Cant find file to read.")
    except Exception as e:
        raise e
    
if __name__ == "__main__":
    print("=" * 60)

    in_csv = "data/samples/people.csv"
    out_xlsx = "data/out/people_from_csv.xlsx"

    try:
        print("Test 3: CSV -> XLSX")
        
        with open(in_csv, "r", encoding="utf-8") as f:
            initial_csv_len = len(f.readlines()) - 1

        csv_to_xlsx(in_csv, out_xlsx)

        wb = openpyxl.load_workbook(out_xlsx)
        ws = wb["Sheet1"]
        final_xlsx_len = ws.max_row - 1

        print(f" -> Zapisi CSV: {initial_csv_len}")
        print(f" -> Zapisi XLSX: {final_xlsx_len}")

        assert initial_csv_len == final_xlsx_len, "Erorr"
        print(" Success")
        print("="*60)
        
    except Exception as e:
        print(f"\nErorr: {e}")